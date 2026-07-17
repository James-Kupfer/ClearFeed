"""Generic profile-driven export: query ContentRecords and save to Excel.

A profile YAML drives all behavior — the query, output path, and filename.
No hardcoded logic for any specific export type.

Usage:
    clearfeed.bat export profiles\\job_export.yaml
    clearfeed.bat export profiles\\some_other_export.yaml

Profile schema (YAML):
    kind: export
    name: job_export
    sql: |
      SELECT cr.subject, cr.sender, cr.processed_at, cr.summary,
             cr.body_text, cr.linked_article_url
      FROM ContentRecords cr
      WHERE cr.enrichment_status = 'complete'
        AND cr.processed_at >= (now() AT TIME ZONE 'America/Chicago') - interval '48 hours'
        AND EXISTS (...)
      ORDER BY cr.processed_at DESC LIMIT 500
    output_path: "O:\\path\\to\\folder"
    filename_template: job_export_{timestamp}.xlsx  # optional
"""

import argparse
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: F401 — side effect: sets up paths for downstream imports
from typing import Any

from db import DbError, query_sql
from utils import purge_old_logs, read_yaml_profile, setup_logging

log = logging.getLogger(__name__)


def _load_profile(profile_path: str) -> dict:
    """Load and minimally validate the export profile YAML."""
    profile = read_yaml_profile(Path(profile_path))

    kind = profile.get("kind")
    if kind is not None and kind != "export":
        raise ValueError(f"Expected kind: export, got {kind!r}")

    for old_key in ("filter", "columns", "window_hours", "limit"):
        if old_key in profile:
            raise ValueError(
                f"Export profile key {old_key!r} is no longer supported. "
                "Replace the filter/columns/window_hours/limit block with a raw 'sql:' statement. "
                "See README for the new profile schema."
            )

    required = ("name", "sql", "output_path")
    missing = [k for k in required if k not in profile]
    if missing:
        raise ValueError(f"Profile missing required keys: {missing}")

    return profile


def _save_excel(
    col_names: list[str],
    rows: list[list[Any]],
    output_path: str,
    filename_template: str,
    profile_name: str,
) -> Path:
    """Write rows to a timestamped Excel file. Returns the output path."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError(
            "openpyxl is required for Excel export. Install it: pip install openpyxl"
        )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = filename_template.replace("{timestamp}", timestamp).replace("{profile}", profile_name)
    out_dir = Path(output_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = profile_name[:31]  # Excel sheet name limit

    # Header row
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=False)

    for col_idx, name in enumerate(col_names, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 18

    # Data rows
    data_font = Font(name="Calibri", size=10)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            # Normalize newlines so Excel displays cleanly
            if isinstance(value, str):
                value = value.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font

    # Auto-width (approximate: cap at 80 chars)
    for col_idx, name in enumerate(col_names, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(name)
        for row in rows:
            val = row[col_idx - 1]
            if val is not None:
                cell_len = len(str(val))
                max_len = max(max_len, min(cell_len, 80))
        ws.column_dimensions[col_letter].width = max_len + 2

    wb.save(dest)
    return dest


def run_export(profile_path: str) -> None:
    """Load profile, query DB, write Excel. Core logic (no sys.exit)."""
    log.info("Loading profile: %s", profile_path)
    profile = _load_profile(profile_path)
    profile_name: str = profile["name"]

    log.info("Profile '%s': executing SQL", profile_name)
    col_names, row_dicts = query_sql(profile["sql"], with_columns=True)
    rows = [[r[c] for c in col_names] for r in row_dicts]
    log.info("Query returned %d row(s)", len(rows))

    if not rows:
        log.info("No records matched — no Excel file written.")
        return

    filename_template: str = profile.get("filename_template", "{profile}_{timestamp}.xlsx")
    output_path: str = profile["output_path"]
    dest = _save_excel(col_names, rows, output_path, filename_template, profile_name)
    log.info("Saved: %s (%d rows)", dest, len(rows))


def main() -> None:
    """Entry point: parse args, run export."""
    setup_logging("export")
    purge_old_logs()

    parser = argparse.ArgumentParser(
        description="ClearFeed export — query ContentRecords and save to Excel per profile."
    )
    parser.add_argument(
        "profile",
        help="Path to the export profile YAML (e.g. profiles\\job_export.yaml)",
    )
    args = parser.parse_args()

    try:
        run_export(args.profile)
    except (FileNotFoundError, ValueError) as exc:
        log.error("Profile error: %s", exc)
        sys.exit(1)
    except DbError as exc:
        log.error("Database error: %s", exc)
        sys.exit(1)
    except RuntimeError as exc:
        log.error("%s", exc)
        sys.exit(1)


if __name__ == "__main__":
    main()
